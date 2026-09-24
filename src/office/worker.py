"""Bounded document operations inside KK Code's offline, immutable tool image.

No user Python/shell is accepted. Originals live in a read-only inputs mount;
outputs are new files and are rendered/reopened before a success receipt.
"""
import hashlib
import importlib.metadata
import io
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

from defusedxml import ElementTree as SafeXML
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches as SlideInches, Pt as SlidePt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import Text
from pypdf.annotations import Link
import pypdfium2 as pdfium
from urllib.parse import urlsplit
from PIL import Image, ImageStat

ROOT = Path("/workspace")
INPUT = ROOT / "inputs"
OUTPUT = ROOT / "outputs"
MAX_PAGES = 40
MAX_TEXT = 200_000
MAX_CELLS = 20_000
FORMULAS = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "IF", "IFERROR", "ROUND", "ROUNDUP", "ROUNDDOWN", "ABS", "AND", "OR", "NOT", "CONCAT", "CONCATENATE", "LEN", "LEFT", "RIGHT", "MID", "TRIM", "UPPER", "LOWER", "SUMIF", "COUNTIF", "SUMIFS", "COUNTIFS", "VLOOKUP", "INDEX", "MATCH", "DATE", "YEAR", "MONTH", "DAY"}


class OfficeError(Exception):
    def __init__(self, code, message, details=None):
        super().__init__(message)
        self.code, self.details = code, details


def reject(message, code="office_unsupported", details=None):
    raise OfficeError(code, message, details)


def bounded_text(value, limit=MAX_TEXT):
    if not isinstance(value, str) or len(value) > limit or "\x00" in value:
        reject("文本类型无效或超过处理上限", "office_input")
    return value


def sequence(value, limit, label):
    if not isinstance(value, list) or len(value) > limit:
        reject(f"{label}必须为不超过 {limit} 项的列表", "office_input")
    return value


def input_path(value):
    if not isinstance(value, str) or not re.fullmatch(r"inputs/[A-Za-z0-9_.-]+", value):
        reject("输入必须来自宿主复制的文件列表", "office_scope")
    target = (ROOT / value).resolve(strict=True)
    if target.parent != INPUT or not target.is_file() or target.stat().st_nlink != 1:
        reject("输入路径不是独立的普通文件", "office_scope")
    if target.stat().st_size > 128 * 1024 * 1024:
        reject("输入文件超过 128 MiB", "office_limit")
    return target


def output_path(value, extension=None):
    if not isinstance(value, str) or not 1 <= len(value) <= 180 or value.startswith(".") or re.search(r'[\\/:*?"<>|\x00-\x1f\x7f]', value) or value.endswith((" ", ".")) or re.match(r"^(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\.|$)", value, re.I):
        reject("输出名称无效", "office_scope")
    target = OUTPUT / value
    if extension and target.suffix.lower() != extension:
        reject("输出扩展名与格式不一致", "office_input")
    if target.exists():
        reject("不覆盖已存在的输出文件", "office_conflict")
    return target


def run(argv, timeout=60):
    try:
        result = subprocess.run(argv, capture_output=True, text=True, timeout=timeout, check=False,
                                env={"PATH": "/usr/local/bin:/usr/bin:/bin", "HOME": "/tmp", "TMPDIR": "/tmp", "LANG": "C.UTF-8", "LC_ALL": "C.UTF-8"})
    except subprocess.TimeoutExpired:
        reject("文档处理超时，未将产物标记为完成", "office_timeout")
    if result.returncode != 0:
        reject(f"文档工具执行失败：{Path(argv[0]).name}", "office_execution", {"exitCode": result.returncode})
    return result.stdout


def inspect_package(filename, depth=0):
    if depth > 1: reject("Office 嵌套对象层数超过上限")
    try:
        archive = zipfile.ZipFile(filename)
    except zipfile.BadZipFile:
        reject("文件不是有效的 Office Open XML，旧格式或加密文档暂不支持")
    with archive:
        entries = archive.infolist()
        if len(entries) > 20_000 or sum(item.file_size for item in entries) > 256 * 1024 * 1024:
            reject("Office 压缩包超过安全解压额度", "office_limit")
        issues = []
        for entry in entries:
            name = entry.filename.lower()
            if entry.flag_bits & 1 or name.startswith("/") or ".." in Path(name).parts or "\\" in name:
                reject("Office 包包含加密或不安全条目")
            chart_workbook = name.startswith("ppt/embeddings/") and name.endswith(".xlsx") and entry.file_size < 32 * 1024 * 1024
            if chart_workbook:
                inspect_package(io.BytesIO(archive.read(entry)), depth + 1)
            if any(part in name for part in ["vbaproject", "/activex/", "/externallinks/", "/connections", "/querytables/", "/pivot", "/slicers/", "/diagrams/", "_xmlsignatures/"]) or ("/embeddings/" in name and not chart_workbook):
                issues.append(entry.filename)
            if name.startswith("ppt/media/") and Path(name).suffix in {".mp4", ".avi", ".mov", ".wmv", ".mp3", ".wav", ".ogg"}:
                issues.append("unsupported_audio_video")
            if name.startswith("xl/") and any(part in name for part in ["threadedcomments", "comments", "vml", "metadata.xml"]):
                issues.append("unsupported_spreadsheet_extensions")
            if name.endswith((".xml", ".rels")):
                if entry.file_size > 16 * 1024 * 1024:
                    reject("Office XML 单个条目过大", "office_limit")
                data = archive.read(entry)
                try:
                    element = SafeXML.fromstring(data)
                except Exception:
                    reject("Office XML 非法或包含不允许的实体定义")
                for node in element.iter():
                    local = node.tag.rsplit("}", 1)[-1]
                    if local in {"ins", "del", "altChunk", "object", "oleObj"}:
                        issues.append(local)
                    if name.startswith("xl/drawings/") and local in {"sp", "grpSp", "cxnSp"}:
                        issues.append("unsupported_spreadsheet_shapes")
                    if local == "Relationship" and node.attrib.get("TargetMode") == "External":
                        # Clickable citation hyperlinks are data, never fetched.
                        if not node.attrib.get("Type", "").endswith("/hyperlink"):
                            issues.append("external_relationship")
                    if local in {"instrText", "fldSimple"} and re.search(r"DDE|INCLUDETEXT|INCLUDEPICTURE|LINK", "".join(node.itertext()) + str(node.attrib), re.I):
                        issues.append("external_field")
        if issues:
            reject("包含当前工具不能安全保真处理的对象；原文件保持不变", details={"unsupportedObjects": sorted(set(issues))[:30]})


def pdf_reader(filename):
    with Path(filename).open("rb") as handle:
        if b"%PDF-" not in handle.read(1024): reject("输入内容不是有效 PDF，不能只修改文件扩展名")
    reader = PdfReader(filename, strict=True)
    if reader.is_encrypted:
        reject("加密 PDF 暂不支持；不会猜测或移除密码")
    if len(reader.pages) > MAX_PAGES:
        reject(f"PDF 超过本轮完整验收上限 {MAX_PAGES} 页", "office_limit")
    root = reader.trailer["/Root"]
    opening = root.get("/OpenAction")
    if opening is not None:
        opening = opening.get_object()
        benign_view = isinstance(opening, list) and len(opening) >= 2 and str(opening[1]) in {"/XYZ", "/Fit", "/FitB", "/FitH", "/FitV", "/FitBH", "/FitBV", "/FitR"}
        benign_jump = hasattr(opening, "get") and opening.get("/S") == "/GoTo"
        if not benign_view and not benign_jump: reject("PDF 包含不受支持的打开动作")
    if root.get("/AA") or root.get("/Perms"):
        reject("带自动动作、数字签名或特殊权限的 PDF 暂不支持")
    acroform = root.get("/AcroForm")
    if acroform and acroform.get_object().get("/XFA"):
        reject("动态 XFA 表单暂不支持")
    for field in (reader.get_fields() or {}).values():
        if field.get("/FT") == "/Sig": reject("数字签名 PDF 不做修改或重新导出")
    names = root.get("/Names")
    if names and any(key in names.get_object() for key in ["/JavaScript", "/EmbeddedFiles"]):
        reject("带脚本或嵌入文件的 PDF 暂不支持")
    for page in reader.pages:
        if page.get("/AA"):
            reject("带自动动作的 PDF 页面暂不支持")
        for annotation in page.get("/Annots", []):
            item = annotation.get_object()
            if item.get("/FT") == "/Sig" or item.get("/AA"):
                reject("带签名或动态动作的 PDF 暂不支持")
            action = item.get("/A")
            if action and action.get_object().get("/S") not in [None, "/URI", "/GoTo"]:
                reject("带可执行动作的 PDF 暂不支持")
    return reader


def image_input(filename):
    try:
        with Image.open(filename) as image:
            if image.format not in {"PNG", "JPEG"} or getattr(image, "n_frames", 1) != 1 or image.width * image.height > 40_000_000:
                reject("文档图片只支持不超过 4000 万像素的静态 PNG/JPEG")
            image.verify()
    except OfficeError:
        raise
    except Exception:
        reject("图片内容不是可解码的 PNG/JPEG；SVG/PDF 需要先明确转换为位图")
    return filename


def formula(value):
    value = bounded_text(value, 4096)
    if not value.startswith("=") or any(part in value for part in ["[", "]", "|", "http:", "https:", "file:"]):
        reject("公式无效或包含外部引用")
    names = re.findall(r"([A-Za-z_][A-Za-z0-9_.]*)\s*\(", value)
    unsupported = [name for name in names if name.upper() not in FORMULAS]
    if unsupported:
        reject("公式不在已验收的本地计算子集中", details={"functions": unsupported})
    return value


def libreoffice(source, target_format, folder):
    profile = Path(tempfile.mkdtemp(prefix="kk-office-profile-"))
    (profile / "user").mkdir()
    (profile / "user/registrymodifications.xcu").write_text('''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>''')
    run(["soffice", "--headless", "--nologo", "--nodefault", "--nolockcheck", "--norestore",
         f"-env:UserInstallation={profile.as_uri()}", "--convert-to", target_format, "--outdir", str(folder), str(source)], 90)
    result = folder / (source.stem + "." + target_format.split(":", 1)[0])
    if not result.is_file() or result.stat().st_size == 0:
        reject("LibreOffice 未生成有效产物，不能将退出码 0 当作转换成功", "office_validation")
    return result


def recalculate(filename):
    folder = Path(tempfile.mkdtemp(prefix="kk-office-calc-"))
    converted = libreoffice(filename, "xlsx:Calc MS Excel 2007 XML", folder)
    shutil.copyfile(converted, filename)
    formulas = load_workbook(filename, data_only=False, keep_links=False)
    values = load_workbook(filename, data_only=True, keep_links=False)
    checked = []
    for sheet in formulas:
        if sheet.max_row * sheet.max_column > MAX_CELLS:
            reject("工作表超过完整验收单元格上限", "office_limit")
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    formula(cell.value)
                    calculated = values[sheet.title][cell.coordinate]
                    if calculated.data_type == "e" or calculated.value is None:
                        reject("公式没有得到有效重算结果", "office_validation", {"sheet": sheet.title, "cell": cell.coordinate, "value": calculated.value})
                    checked.append({"sheet": sheet.title, "cell": cell.coordinate, "formula": cell.value, "value": calculated.value})
    return {"engine": "LibreOffice Calc", "recalculated": True, "formulaCount": len(checked), "formulas": checked[:1000]}


def preserve_docx_links(source, pdf):
    """Some renderer builds preserve link text but omit URI annotations.

    Restore only provably unique, visible link labels using PDFium's actual
    rendered character-range boxes. Ambiguous labels fail closed; never guess coordinates
    or claim fidelity from a relationship XML entry alone.
    """
    document = Document(source)
    expected = {}
    for link in document._element.iter(qn("w:hyperlink")):
        relation = link.get(qn("r:id"))
        if not relation:
            continue
        target = document.part.rels[relation].target_ref
        parsed = urlsplit(target)
        if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
            reject("仅支持无凭据的 HTTP(S) 文档超链接", "office_unsupported")
        label = "".join(node.text or "" for node in link.iter(qn("w:t")))
        if not label.strip(): reject("不能验证空文本超链接", "office_validation")
        if len(label) > 2000 or len(expected) >= 500: reject("文档超链接数量或文字长度超过完整定位上限", "office_limit")
        previous = expected.get(label)
        if previous and previous["url"] != target: reject("相同超链接文字对应不同地址，无法安全定位 PDF 链接", "office_validation")
        expected.setdefault(label, {"url": target, "count": 0})["count"] += 1
    if not expected:
        return {"hyperlinksVerified": None, "linkCount": 0, "repairedLinkAnnotations": 0}
    reader = pdf_reader(pdf)
    def uri(ref):
        action = ref.get_object().get("/A")
        return str(action.get_object().get("/URI", "")) if action is not None else ""
    existing = [uri(ref) for page in reader.pages for ref in page.get("/Annots", [])]
    missing = {label: item for label, item in expected.items() if existing.count(item["url"]) < item["count"]}
    repaired = 0
    if missing:
        writer = PdfWriter(); writer.clone_document_from_reader(reader)
        with pdfium.PdfDocument(str(pdf)) as layout:
            if len(layout) != len(reader.pages): reject("PDF 链接定位页数与实际文档不一致", "office_validation")
            for label, item in missing.items():
                if existing.count(item["url"]): reject("渲染器只保留了部分同址链接，需人工核验后再交付", "office_validation")
                matches = []
                for page_index in range(len(layout)):
                    page = layout[page_index]; textpage = page.get_textpage()
                    search = textpage.search(label, match_case=True)
                    try:
                        while True:
                            match = search.get_next()
                            if match is None: break
                            count = textpage.count_rects(*match)
                            if not 0 < count <= 100: reject("PDF 超链接范围过于复杂", "office_validation")
                            rects = [textpage.get_rect(index) for index in range(count)]
                            matches.append((page_index, page.get_width(), page.get_height(), rects))
                    finally:
                        search.close(); textpage.close(); page.close()
                if len(matches) != item["count"]:
                    reject("PDF 超链接文字缺失或存在歧义；未将只有文本的链接标为可点击", "office_validation")
                for page_index, width, height, rects in matches:
                    for rect in rects:
                        if not all(math.isfinite(value) for value in rect) or not 0 <= rect[0] < rect[2] <= width or not 0 <= rect[1] < rect[3] <= height:
                            reject("PDF 超链接位于有效页面边界之外", "office_validation")
                        writer.add_annotation(page_number=page_index, annotation=Link(rect=rect, url=item["url"]))
                        repaired += 1
        temporary = pdf.with_suffix(".links.tmp")
        with temporary.open("xb") as handle: writer.write(handle)
        temporary.replace(pdf)
    checked = pdf_reader(pdf)
    actual = [uri(ref) for page in checked.pages for ref in page.get("/Annots", [])]
    if any(actual.count(item["url"]) < item["count"] for item in expected.values()): reject("PDF 超链接回读验证失败", "office_validation")
    return {"hyperlinksVerified": True, "linkCount": sum(item["count"] for item in expected.values()), "repairedLinkAnnotations": repaired}


def render_pdf(filename, expected_text=None, link_source=None):
    if filename.suffix.lower() == ".pdf":
        pdf = filename
    elif filename.suffix.lower() in {".docx", ".xlsx", ".pptx"}:
        pdf = libreoffice(filename, "pdf", OUTPUT)
    else:
        reject("此格式不能直接生成页面预览")
    link_validation = preserve_docx_links(link_source or filename, pdf) if link_source or filename.suffix.lower() == ".docx" else {}
    reader = pdf_reader(pdf)
    if not reader.pages:
        reject("渲染结果没有页面", "office_validation")
    extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
    normal = lambda value: re.sub(r"\s+", "", value)
    missing = [value for value in expected_text or [] if value and normal(value) not in normal(extracted)]
    if missing:
        reject("PDF 渲染后缺少预期文字，可能存在缺字或内容丢失", "office_validation", {"missingText": missing[:20]})
    prefix = OUTPUT / (filename.stem + "-page")
    run(["pdftoppm", "-png", "-scale-to", "1600", str(pdf), str(prefix)], 90)
    previews = sorted(OUTPUT.glob(prefix.name + "-*.png"))
    if len(previews) != len(reader.pages):
        reject("页面预览数量与 PDF 不一致", "office_validation")
    pages = []
    for preview in previews:
        with Image.open(preview) as image:
            image.load()
            if min(image.size) < 100:
                reject("页面预览尺寸无效", "office_validation")
            variance = sum(ImageStat.Stat(image.convert("RGB")).var)
            pages.append({"preview": preview.name, "width": image.width, "height": image.height, "nonBlank": variance > 0.01})
    return {"pdf": pdf.name, "pageCount": len(reader.pages), "pages": pages, **link_validation,
            "textVerified": not missing if any(expected_text or []) else None, "layout": "rendered_for_review", "note": "已实际渲染；仅在提供预期文字时进行文字核验。复杂排版仍需查看预览，不宣称完全自动视觉审校。"}


def create_docx(spec, filename, inputs):
    if not spec.get("title") and not spec.get("blocks"): reject("文档缺少标题或正文内容", "office_input")
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"; normal.font.size = Pt(11)
    expected = []
    if spec.get("title"):
        title = bounded_text(spec["title"], 300); document.add_heading(title, 0); expected.append(title)
    for block in sequence(spec.get("blocks", []), 1000, "文档段落"):
        kind = block.get("type")
        if kind in {"paragraph", "heading"}:
            text = bounded_text(block.get("text", ""), 20000); expected.append(text)
            if kind == "heading":
                level = block.get("level", 1)
                if not isinstance(level, int) or not 1 <= level <= 6: reject("标题层级应为 1–6", "office_input")
                document.add_heading(text, level)
            else:
                paragraph = document.add_paragraph(style="List Bullet" if block.get("style") == "bullet" else "Normal")
                cursor = 0
                for match in re.finditer(r"\[([^\]]+)\]\((https?://[^\s)]+)\)", text):
                    paragraph.add_run(text[cursor:match.start()])
                    relation = paragraph.part.relate_to(match[2], RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
                    hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("r:id"), relation); hyperlink.set(qn("w:history"), "1")
                    run_element = OxmlElement("w:r"); text_element = OxmlElement("w:t"); text_element.text = match[1]
                    properties = OxmlElement("w:rPr")
                    color = OxmlElement("w:color"); color.set(qn("w:val"), "1C5AAA"); properties.append(color)
                    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); properties.append(underline)
                    run_element.append(properties)
                    run_element.append(text_element); hyperlink.append(run_element); paragraph._p.append(hyperlink)
                    cursor = match.end()
                paragraph.add_run(text[cursor:])
                expected[-1] = re.sub(r"\[([^\]]+)\]\((https?://[^\s)]+)\)", r"\1", text)
        elif kind == "table":
            rows = sequence(block.get("rows"), 500, "表格行")
            if not rows: reject("表格不能为空", "office_input")
            columns = len(sequence(rows[0], 20, "表格列"))
            if not columns: reject("表格不能为空", "office_input")
            table = document.add_table(rows=len(rows), cols=columns); table.style = "Table Grid"
            for index, row in enumerate(rows):
                if len(row) != columns: reject("表格各行列数必须相同", "office_input")
                for col, value in enumerate(row):
                    value = bounded_text(str(value), 4000); table.cell(index, col).text = value; expected.append(value)
        elif kind == "image":
            index = block.get("input")
            if not isinstance(index, int) or not 0 <= index < len(inputs): reject("图片输入序号无效", "office_input")
            document.add_picture(str(image_input(inputs[index])), width=Inches(min(6, max(0.5, float(block.get("widthInches", 5))))))
        else: reject("不支持的 Word 内容块", "office_input")
    document.save(filename)
    reopened = Document(filename)
    return expected, {"paragraphs": len(reopened.paragraphs), "tables": len(reopened.tables)}


def set_cell(cell, value):
    if isinstance(value, dict):
        if set(value) != {"formula"}: reject("单元格对象只支持显式 formula", "office_input")
        cell.value = formula(value["formula"])
    elif isinstance(value, (str, int, float, bool)) or value is None:
        cell.value = value
        if isinstance(value, str):
            bounded_text(value, 32000)
            cell.data_type = "s"  # Data beginning '=' is text, never implicit code.
    else: reject("单元格值类型不支持", "office_input")


def create_xlsx(spec, filename):
    book = Workbook(); book.remove(book.active)
    sheets = sequence(spec.get("sheets", []), 30, "工作表")
    if not sheets: reject("至少需要一个工作表", "office_input")
    for sheet_spec in sheets:
        name = bounded_text(sheet_spec.get("name", "Sheet"), 31)
        if not name or re.search(r"[\\/*?:\[\]]", name) or name in book.sheetnames: reject("工作表名称无效或重复", "office_input")
        sheet = book.create_sheet(name)
        rows = sequence(sheet_spec.get("rows", []), 2000, "工作表行")
        if sum(len(row) for row in rows) > MAX_CELLS: reject("工作表超过单元格上限", "office_limit")
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(sequence(row, 200, "工作表列"), 1): set_cell(sheet.cell(r, c), value)
        sheet.freeze_panes = "A2"
        for cell in sheet[1]: cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="E4E7DD")
        for column in sheet.columns: sheet.column_dimensions[column[0].column_letter].width = 20
        chart_spec = sheet_spec.get("chart")
        if chart_spec:
            chart = LineChart() if chart_spec.get("type") == "line" else BarChart()
            if chart_spec.get("type") not in {"line", "bar"}: reject("图表只支持 line/bar", "office_input")
            chart.title = bounded_text(chart_spec.get("title", ""), 200)
            chart.add_data(Reference(sheet, min_col=2, max_col=sheet.max_column, min_row=1, max_row=sheet.max_row), titles_from_data=True)
            chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row))
            sheet.add_chart(chart, "A" + str(sheet.max_row + 3))
    book.save(filename)
    return [], {"sheets": book.sheetnames, **recalculate(filename)}


def create_pptx(spec, filename, inputs):
    presentation = Presentation(); presentation.slide_width = SlideInches(13.333); presentation.slide_height = SlideInches(7.5)
    expected = []
    slides = sequence(spec.get("slides", []), MAX_PAGES, "幻灯片")
    if not slides: reject("至少需要一张幻灯片", "office_input")
    for item in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = bounded_text(item.get("title", ""), 160)
        title_box = slide.shapes.add_textbox(SlideInches(0.6), SlideInches(0.35), SlideInches(12.1), SlideInches(0.8))
        title_box.text_frame.text = title; expected.append(title)
        title_box.text_frame.paragraphs[0].font.size = SlidePt(30)
        body = sequence(item.get("body", []), 8, "幻灯片正文")
        if sum(len(str(line)) for line in body) > 900: reject("单页正文过长，无法保证不裁切；请拆分幻灯片", "office_layout")
        if body:
            box = slide.shapes.add_textbox(SlideInches(0.8), SlideInches(1.6), SlideInches(6 if item.get("image") is not None or item.get("chart") else 11.5), SlideInches(5.2))
            for index, line in enumerate(body):
                line = bounded_text(line, 240); expected.append(line)
                paragraph = box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
                paragraph.text = line; paragraph.font.size = SlidePt(21); paragraph.space_after = SlidePt(12)
        if item.get("image") is not None:
            index = item["image"]
            if not isinstance(index, int) or not 0 <= index < len(inputs): reject("图片输入序号无效", "office_input")
            with Image.open(image_input(inputs[index])) as image: ratio = image.height / image.width
            width = min(5.2, 4.8 / ratio)
            slide.shapes.add_picture(str(inputs[index]), SlideInches(7.3), SlideInches(1.7), width=SlideInches(width))
        if item.get("chart"):
            chart = item["chart"]
            if chart.get("type") not in {"bar", "line"}: reject("图表只支持 line/bar", "office_input")
            data = CategoryChartData(); data.categories = sequence(chart.get("categories"), 30, "图表类别")
            for series in sequence(chart.get("series", []), 10, "图表数据系列"):
                values = sequence(series.get("values"), 30, "图表数值")
                if len(values) != len(data.categories): reject("图表数值与类别数量不一致", "office_input")
                data.add_series(bounded_text(series.get("name", ""), 100), values)
            slide.shapes.add_chart(XL_CHART_TYPE.LINE if chart["type"] == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED,
                                   SlideInches(7), SlideInches(1.7), SlideInches(5.5), SlideInches(4.8), data)
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > presentation.slide_width or shape.top + shape.height > presentation.slide_height:
                reject("幻灯片对象超出页面边界", "office_layout")
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.font.name = "Noto Sans CJK SC"
    presentation.save(filename)
    return expected, {"slides": len(Presentation(filename).slides), "boundsChecked": True}


def inspect_document(filename):
    kind = filename.suffix.lower()
    if kind in {".docx", ".xlsx", ".pptx"}: inspect_package(filename)
    if kind == ".docx":
        doc = Document(filename)
        return {"format": "docx", "paragraphs": [p.text for p in doc.paragraphs], "tables": [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]}
    if kind == ".xlsx":
        book = load_workbook(filename, data_only=False, keep_links=False, rich_text=True)
        sheets = []
        for sheet in book:
            if sheet.max_row * sheet.max_column > MAX_CELLS: reject("工作表超过完整读取上限", "office_limit")
            sheets.append({"name": sheet.title, "rows": [[cell.value for cell in row] for row in sheet]})
        return {"format": "xlsx", "sheets": sheets, "formulaValues": "inspect retains formulas; recalculate uses LibreOffice"}
    if kind == ".pptx":
        presentation = Presentation(filename)
        return {"format": "pptx", "slides": [[shape.text for shape in slide.shapes if shape.has_text_frame] for slide in presentation.slides]}
    if kind == ".pdf":
        reader = pdf_reader(filename)
        pages = [{"page": index + 1, "text": page.extract_text() or ""} for index, page in enumerate(reader.pages)]
        return {"format": "pdf", "pageCount": len(pages), "pages": pages, "needsOcr": any(not p["text"].strip() for p in pages),
                "fields": {name: {"type": field.get("/FT"), "value": field.get("/V")} for name, field in (reader.get_fields() or {}).items()}}
    if kind == ".md": return {"format": "markdown", "text": bounded_text(filename.read_text(encoding="utf8"))}
    reject("不支持此输入格式")


def replacements(paragraphs, changes):
    for change in sequence(changes, 100, "文本替换"):
        find = bounded_text(change.get("find", ""), 5000); replacement = bounded_text(change.get("replace", ""), 5000)
        if not find: reject("查找文本不能为空", "office_input")
        matches = [(paragraph, run) for paragraph in paragraphs for run in paragraph.runs if find in run.text]
        if len(matches) != 1: reject("替换必须唯一命中一个文本 run；跨格式或多处文本请先人工定位", "office_unsupported")
        matches[0][1].text = matches[0][1].text.replace(find, replacement)


def edit_document(source, filename, changes):
    if source.suffix == ".md":
        content = bounded_text(source.read_text(encoding="utf8"))
        for change in sequence(changes.get("replace", []), 100, "文本替换"):
            find = bounded_text(change.get("find", "")); replacement = bounded_text(change.get("replace", ""))
            if not find or content.count(find) != 1: reject("Markdown 替换必须唯一命中", "office_input")
            content = content.replace(find, replacement)
        filename.write_text(content, encoding="utf8")
        return [], {"roundTrip": filename.read_text(encoding="utf8") == content, "originalPreserved": True}
    inspect_package(source)
    if source.suffix == ".docx":
        doc = Document(source)
        paragraphs = list(doc.paragraphs) + [p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
        replacements(paragraphs, changes.get("replace", [])); doc.save(filename)
        return [p.text for p in doc.paragraphs if p.text], {"edited": "unique text runs", "originalPreserved": True}
    if source.suffix == ".xlsx":
        book = load_workbook(source, keep_links=False, rich_text=True)
        for change in sequence(changes.get("cells", []), MAX_CELLS, "单元格修改"):
            if change.get("sheet") not in book.sheetnames or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,5}", change.get("cell", "")): reject("单元格地址无效", "office_input")
            set_cell(book[change["sheet"]][change["cell"]], change.get("value"))
        book.save(filename)
        return [], {"originalPreserved": True, **recalculate(filename)}
    if source.suffix == ".pptx":
        presentation = Presentation(source)
        paragraphs = [paragraph for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame for paragraph in shape.text_frame.paragraphs]
        replacements(paragraphs, changes.get("replace", [])); presentation.save(filename)
        return [p.text for p in paragraphs if p.text], {"edited": "unique text runs", "originalPreserved": True}
    reject("此格式不支持保真局部编辑")


def capabilities():
    versions = {name: importlib.metadata.version(name) for name in ["python-docx", "openpyxl", "python-pptx", "pypdf", "pypdfium2", "Pillow", "defusedxml"]}
    return {"protocol": 1, "versions": versions, "libreoffice": run(["soffice", "--version"]).strip(),
            "operations": ["capabilities", "inspect", "create", "edit", "render", "merge_pdf", "select_pdf_pages", "fill_pdf", "annotate_pdf", "ocr"],
            "formats": ["docx", "xlsx", "pptx", "pdf", "md"], "ocrLanguages": run(["tesseract", "--list-langs"]).splitlines()[1:],
            "limits": {"pages": MAX_PAGES, "cellsPerSheet": MAX_CELLS, "inputBytes": 128 * 1024 * 1024},
            "unsupported": ["legacy doc/xls/ppt", "macros", "encrypted/signed documents", "embedded OLE/ActiveX/SmartArt", "tracked changes", "arbitrary PDF rewrite"]}


def execute(request):
    if not isinstance(request, dict) or len(json.dumps(request)) > 2 * 1024 * 1024: reject("请求格式或大小无效", "office_input")
    operation = request.get("operation")
    if operation == "capabilities": return {"capabilities": capabilities()}
    OUTPUT.mkdir(exist_ok=True)
    inputs = [input_path(value) for value in sequence(request.get("inputs", []), 30, "输入文件")]
    source = inputs[0] if inputs else None
    if operation == "inspect":
        if not source: reject("缺少输入文件", "office_input")
        content = inspect_document(source)
        if len(json.dumps(content, ensure_ascii=False, default=str)) > 2 * 1024 * 1024: reject("读取结果超过上限，不返回截断的完整声明", "office_limit")
        return {"content": content, "complete": True}
    expected, structural, rendered, link_source = [], {}, None, None
    if operation == "create":
        kind = request.get("format")
        if kind not in {"docx", "xlsx", "pptx", "pdf", "md"}: reject("创建格式无效", "office_input")
        filename = output_path(request.get("filename", f"document.{kind}"), "." + kind)
        spec = request.get("spec", {})
        if not isinstance(spec, dict): reject("spec 必须是对应文档格式的对象", "office_input")
        if kind in {"docx", "pdf"}:
            intermediate = filename if kind == "docx" else OUTPUT / (filename.stem + ".docx")
            expected, structural = create_docx(spec, intermediate, inputs)
            if kind == "pdf": link_source = intermediate; filename = libreoffice(intermediate, "pdf", OUTPUT)
        elif kind == "xlsx": expected, structural = create_xlsx(spec, filename)
        elif kind == "pptx": expected, structural = create_pptx(spec, filename, inputs)
        else:
            text = bounded_text(spec.get("text", "")); filename.write_text(text, encoding="utf8")
            structural = {"roundTrip": filename.read_text(encoding="utf8") == text, "characters": len(text)}
    elif operation == "edit":
        if not source: reject("缺少输入文件", "office_input")
        filename = output_path(request.get("filename", "edited" + source.suffix), source.suffix)
        changes = request.get("changes", {})
        if not isinstance(changes, dict): reject("changes 必须是编辑操作对象", "office_input")
        expected, structural = edit_document(source, filename, changes)
    elif operation == "render":
        if not source: reject("缺少输入文件", "office_input")
        inspect_document(source)
        if source.suffix == ".md":
            blocks, fenced = [], False
            for line in bounded_text(source.read_text(encoding="utf8")).splitlines():
                if line.startswith("```"): fenced = not fenced; continue
                heading = re.match(r"^(#{1,6})\s+(.+)$", line) if not fenced else None
                if heading: blocks.append({"type": "heading", "text": heading[2], "level": len(heading[1])})
                elif re.match(r"^[-*]\s+", line) and not fenced: blocks.append({"type": "paragraph", "style": "bullet", "text": line[2:]})
                elif line: blocks.append({"type": "paragraph", "text": line})
            filename = OUTPUT / "document.docx"
            expected, structural = create_docx({"blocks": blocks}, filename, inputs)
            structural["markdownSubset"] = "headings, paragraphs, bullets, code text, clickable HTTP(S) links; advanced layout remains literal"
        else:
            filename = OUTPUT / ("document" + source.suffix); shutil.copyfile(source, filename)
    elif operation in {"merge_pdf", "select_pdf_pages"}:
        if not inputs or any(item.suffix != ".pdf" for item in inputs): reject("此操作仅接受 PDF 输入", "office_input")
        filename = output_path(request.get("filename", "document.pdf"), ".pdf")
        writer = PdfWriter()
        if operation == "merge_pdf":
            seen_fields = set()
            for item in inputs:
                reader = pdf_reader(item); fields = set((reader.get_fields() or {}).keys())
                if seen_fields.intersection(fields): reject("合并输入含同名表单字段，不能静默覆盖；请先区分字段名称")
                seen_fields.update(fields); writer.append(reader)
        else:
            reader = pdf_reader(source)
            pages = sequence(request.get("pages", []), MAX_PAGES, "页面选择")
            for page in pages:
                if not isinstance(page, int) or not 1 <= page <= len(reader.pages): reject("PDF 页码越界", "office_input")
            writer.append(reader, pages=[page - 1 for page in pages])
        if not 0 < len(writer.pages) <= MAX_PAGES: reject("输出页数无效或超过上限", "office_limit")
        with filename.open("xb") as handle: writer.write(handle)
        structural = {"pageCount": len(writer.pages), "originalPreserved": True}
    elif operation in {"fill_pdf", "annotate_pdf"}:
        if not source or source.suffix != ".pdf": reject("此操作仅接受 PDF", "office_input")
        reader = pdf_reader(source); writer = PdfWriter(); writer.append(reader)
        filename = output_path(request.get("filename", "updated.pdf"), ".pdf")
        if operation == "fill_pdf":
            fields = request.get("fields", {})
            available = reader.get_fields() or {}
            if not isinstance(fields, dict) or not 0 < len(fields) <= 100: reject("请提供 1–100 个文本表单字段", "office_input")
            for name, value in fields.items():
                if name not in available or available[name].get("/FT") != "/Tx": reject("仅支持已存在的 AcroForm 文本字段", "office_unsupported")
                bounded_text(value, 4000)
            for page in writer.pages:
                if page.get("/Annots"): writer.update_page_form_field_values(page, fields, auto_regenerate=False)
            with filename.open("xb") as handle: writer.write(handle)
            actual = pdf_reader(filename).get_form_text_fields()
            if any(actual.get(name) != value for name, value in fields.items()): reject("表单字段回读不一致", "office_validation")
            structural = {"filledFields": fields, "valuesVerified": True, "requiresVisualReview": True}
        else:
            annotations = sequence(request.get("annotations", []), 100, "PDF 便笺")
            if not annotations: reject("便笺列表不能为空", "office_input")
            for item in annotations:
                page_number, rect = item.get("page"), item.get("rect")
                if not isinstance(page_number, int) or not 1 <= page_number <= len(writer.pages): reject("便笺页码越界", "office_input")
                if not isinstance(rect, list) or len(rect) != 4 or not all(isinstance(value, (int, float)) for value in rect): reject("便笺坐标必须为四个 PDF point 数值", "office_input")
                box = writer.pages[page_number - 1].mediabox
                if not (float(box.left) <= rect[0] < rect[2] <= float(box.right) and float(box.bottom) <= rect[1] < rect[3] <= float(box.top)): reject("便笺坐标超出页面", "office_layout")
                writer.add_annotation(page_number=page_number - 1, annotation=Text(rect=tuple(rect), text=bounded_text(item.get("text", ""), 4000), open=False))
            with filename.open("xb") as handle: writer.write(handle)
            actual = pdf_reader(filename)
            contents = [str(annotation.get_object().get("/Contents", "")) for page in actual.pages for annotation in page.get("/Annots", [])]
            if any(item["text"] not in contents for item in annotations): reject("便笺内容回读不一致", "office_validation")
            structural = {"notesAdded": len(annotations), "contentsVerified": True, "note": "文本便笺由阅读器点击展开，不是对页面正文的改写。"}
    elif operation == "ocr":
        if not source or source.suffix not in {".pdf", ".png", ".jpg", ".jpeg"}: reject("OCR 仅接受 PDF/PNG/JPEG", "office_input")
        language = request.get("language", "eng+chi_sim")
        if language not in {"eng", "chi_sim", "eng+chi_sim"}: reject("OCR 语言未安装", "office_capability")
        images = [source]
        if source.suffix != ".pdf": image_input(source)
        if source.suffix == ".pdf":
            pdf_reader(source)
            prefix = OUTPUT / "ocr-source"
            run(["pdftoppm", "-png", "-r", "150", str(source), str(prefix)], 90)
            images = sorted(OUTPUT.glob("ocr-source-*.png"))
        texts, confidences = [], []
        for item in images:
            base = OUTPUT / (item.stem + "-ocr")
            run(["tesseract", str(item), str(base), "-l", language, "txt", "tsv"], 60)
            texts.append(base.with_suffix(".txt").read_text(encoding="utf8"))
            lines = base.with_suffix(".tsv").read_text(encoding="utf8").splitlines()
            for line in lines[1:]:
                columns = line.split("\t")
                if len(columns) >= 12 and columns[11].strip() and float(columns[10]) >= 0: confidences.append(float(columns[10]))
        filename = output_path(request.get("filename", "recognized.md"), ".md")
        filename.write_text("\n\n".join(texts), encoding="utf8")
        structural = {"ocr": True, "language": language, "averageConfidence": sum(confidences) / len(confidences) if confidences else None,
                      "requiresReview": True, "warning": "OCR 结果不是原始文本，低置信字符和版面顺序需要人工核对。"}
    else: reject("不支持此文档操作", "office_input")
    if filename.suffix in {".docx", ".xlsx", ".pptx", ".pdf"}:
        inspect_document(filename)
        rendered = render_pdf(filename, expected, link_source)
        if operation == "create" and not any(page["nonBlank"] for page in rendered["pages"]): reject("新建文档渲染后没有可见内容", "office_validation")
    outputs = []
    for item in sorted(OUTPUT.iterdir()):
        if item.is_file():
            outputs.append({"name": item.name, "bytes": item.stat().st_size, "sha256": hashlib.sha256(item.read_bytes()).hexdigest()})
    if sum(item["bytes"] for item in outputs) > 256 * 1024 * 1024: reject("输出超过 256 MiB", "office_limit")
    return {"primary": filename.name, "outputs": outputs, "validation": {"structural": structural, "render": rendered, "originalPreserved": True}}


if __name__ == "__main__":
    try:
        raw = sys.stdin.read(2 * 1024 * 1024 + 1)
        if len(raw) > 2 * 1024 * 1024: reject("请求超过 2 MiB", "office_limit")
        result = execute(json.loads(raw))
        print(json.dumps({"ok": True, **result}, ensure_ascii=False, default=str))
    except OfficeError as error:
        print(json.dumps({"ok": False, "error": {"code": error.code, "message": str(error), "details": error.details}}, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"ok": False, "error": {"code": "office_execution", "message": f"文档处理失败（{type(error).__name__}）；未将产物标记为完成，原件不变。"}}, ensure_ascii=False))
